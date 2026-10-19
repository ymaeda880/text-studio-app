# -*- coding: utf-8 -*-
# /Users/macmini2025/projects/text_studio_project/
# text_studio_app/lib/ai_chat/exports.py
# ============================================================
# 会話全体をWord/PDFへ出力する（出典URLも含める）．
# 生成は利用者がボタンを押したときだけ行う．
# ============================================================
import re
from io import BytesIO
from xml.sax.saxutils import escape


# ============================================================
# 本文と出典を重複させず文書出力用に並べる
# ============================================================
def sections(job):
    yield job['title'], 'title'
    ctx = job.get('doc_context')
    if ctx and ctx.get('text'):
        yield '前提文書', 'heading'
        yield ctx['text'], 'body'
    for i, msg in enumerate(job['messages'], 1):
        yield f"{i}．" + {'user': 'ユーザー', 'assistant': 'AI', 'system': 'システム'}[msg['role']], 'heading'
        yield msg['content'], 'body'
        search = msg.get('search') or {}

        # 回答本文にURL一覧がある場合は，同じ一覧を再出力しない．
        if not search.get('sources_in_text'):
            for n, source in enumerate(search.get('sources', []), 1):
                yield f"出典{n}：{source['title']}\n{source['url']}", 'body' 


# ============================================================
# 日本語の会話全体をWordへ出力する
# ============================================================
def make_docx(job):
    from docx import Document
    from docx.shared import Pt, Cm
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    doc = Document()
    section = doc.sections[0]
    section.top_margin = section.bottom_margin = Cm(2)
    normal = doc.styles['Normal']
    normal.font.name = 'Yu Gothic'
    normal.font.size = Pt(10.5)
    fonts = normal.element.get_or_add_rPr()
    rfonts = OxmlElement('w:rFonts')
    rfonts.set(qn('w:eastAsia'), 'Yu Gothic')
    fonts.append(rfonts)
    for text, kind in sections(job):
        if kind == 'title':
            doc.add_heading(text, 0)
        elif kind == 'heading':
            doc.add_heading(text, 2)
        else:
            for line in text.splitlines():
                heading = re.match(r'^#{1,6}\s+(.+)', line)
                if heading:
                    doc.add_heading(heading.group(1), 3)
                else:
                    doc.add_paragraph(line)
    out = BytesIO()
    doc.save(out)
    return out.getvalue()


# ============================================================
# 日本語の会話全体をPDFへ出力する
# ============================================================
def make_pdf(job):
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.pagesizes import A4
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
    from reportlab.lib.enums import TA_LEFT
    font = 'HeiseiKakuGo-W5'
    pdfmetrics.registerFont(UnicodeCIDFont(font))
    styles = {kind: ParagraphStyle(kind, fontName=font, fontSize=size,
                leading=size * 1.6, wordWrap='CJK', alignment=TA_LEFT,
                splitLongWords=True, spaceAfter=7, keepWithNext=kind != 'body')
              for kind, size in [('title', 18), ('heading', 12), ('body', 10)]}
    out = BytesIO()
    story = []
    for text, kind in sections(job):

        # 本文を行ごとに分け，長い会話でもページ境界で分割できるようにする．
        for line in text.splitlines() or ['']:
            heading = re.match(r'^#{1,6}\s+(.+)', line)
            if kind == 'body' and heading:
                story.append(Paragraph(escape(heading.group(1)), styles['heading']))
            else:
                story.append(Paragraph(escape(line) or '&#160;', styles[kind]))
        story.append(Spacer(1, 5))
    SimpleDocTemplate(out, pagesize=A4, leftMargin=42, rightMargin=42,
                      topMargin=42, bottomMargin=42, title=job['title']).build(story)
    return out.getvalue()
